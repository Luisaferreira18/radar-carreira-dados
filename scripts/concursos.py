"""Baixa e limpa as Tabelas de Autorizações e Provimentos de concursos federais.

Fonte: Ministério da Gestão e da Inovação em Serviços Públicos (MGI), Observatório de
Pessoal - https://www.gov.br/servidor/pt-br/observatorio-de-pessoal-govbr/autorizacoes-e-provimentos-concursos

Cada ano é uma planilha .xlsx com formato um pouco diferente. A limpeza:
  - localiza a linha de cabeçalho e mapeia as colunas pelo nome (que muda entre anos);
  - remove linhas de total, linhas vazias e duplicatas exatas;
  - converte VAGAS para número ("1.000" -> 1000; "CR" = cadastro reserva -> 0);
  - agrupa os >20 textos de "tipo de autorização" em 4 categorias;
  - extrai a data do ato oficial (texto ou link do DOU);
  - padroniza escolaridade (NS/NI/NA) e marca cargos de TI e dados por palavra-chave.

Os links das planilhas são descobertos na própria página oficial (o nome do arquivo
muda a cada atualização). Cada planilha só é aceita se o ano do título bater com o
ano do link — é assim que se detecta que o link "2023" abre, na verdade, a tabela de 2024.

Uso:  python scripts/concursos.py
"""

from __future__ import annotations

import html
import json
import re
import unicodedata
from datetime import date, datetime, timezone

import openpyxl

from comum import DIR_PUBLICO, DIR_RAW, RAIZ, baixar, escrever_se_mudou, salvar_json

PAGINA = "https://www.gov.br/servidor/pt-br/observatorio-de-pessoal-govbr/autorizacoes-e-provimentos-concursos"
GESTAO = "https://www.gov.br/gestao/pt-br/acesso-a-informacao/servidores/concursos-publicos/arquivos-autorizacoes/"
SERVIDOR = "https://www.gov.br/servidor/pt-br/"

# Links conhecidos (usados se a página oficial estiver fora do ar).
PLANILHAS_PADRAO = {
    2016: GESTAO + "atualizacao-de-provimentos_dez_2016.xlsx",
    2017: GESTAO + "planilha-provimentos-autorizados-2017.xlsx",
    2018: GESTAO + "planilha-de-provimentos-autorizados-2018-2.xlsx",
    2019: GESTAO + "provimentos-autorizados-dez-2019-1.xlsx",
    2020: SERVIDOR + "arquivos/provimentosautorizados24dezembro2020_retific_24fevereiro2021.xlsx",
    2021: SERVIDOR + "arquivos/autorizacoes31dezembro2021.xlsx",
    2022: SERVIDOR + "arquivos/autorizacoes1julho2022.xlsx",
    2024: GESTAO + "copy11_of_autorizacoes31dez2024.xlsx",
    2025: SERVIDOR + "observatorio-de-pessoal-govbr/arquivos/autorizacoes28agosto20251.xlsx",
    2026: SERVIDOR + "observatorio-de-pessoal-govbr/arquivos-2026/copy_of_autorizacoes19fevereiro2026.xlsx/@@download/file",
}
ANO_INICIAL = 2016
DIR_CONCURSOS = DIR_RAW / "concursos"
FONTES = DIR_CONCURSOS / "fontes.json"  # ano -> URL de onde cada planilha foi baixada

MESES = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6, "julho": 7,
    "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}
PADRAO_TI = re.compile(
    r"tecnologia da informa|\bti\b|cientista de dados|ciencia de dados|\bdados\b|estatistic|"
    r"informatica|computacao|desenvolvimento de sistemas|analista de sistemas|seguranca da informacao",
)


def sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)).lower()


def limpar_texto(valor: object) -> str:
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor).replace("\xa0", " ")).strip()


def identificar_colunas(cabecalho: tuple) -> dict[str, int]:
    regras = {
        "orgao": lambda h: h.startswith("orgao"),
        "vinculo": lambda h: h.startswith("vinculo"),
        "cargos": lambda h: h.startswith("cargo"),
        "escolaridade": lambda h: h.startswith("escolaridade"),
        "vagas": lambda h: h.startswith("vagas"),
        "ato": lambda h: h.startswith("publicacao") or h.startswith("ato oficial"),
        "link": lambda h: h.startswith("link"),
        "area": lambda h: h.startswith("area"),
        "tipo": lambda h: h.startswith("tipo"),
        "inicio": lambda h: h.startswith("data provimento"),
    }
    colunas: dict[str, int] = {}
    for i, valor in enumerate(cabecalho):
        h = sem_acento(limpar_texto(valor))
        for nome, regra in regras.items():
            if nome not in colunas and h and regra(h):
                colunas[nome] = i
    return colunas


def classificar_tipo(texto: str) -> str:
    t = sem_acento(texto)
    if "tempor" in t or re.search(r"\bctu\b", t):
        return "Contratação temporária"
    if "concurso" in t:
        return "Concurso público"
    if "judicial" in t:
        return "Decisão judicial"
    return "Nomeação de aprovados"  # provimento originário/adicional/excepcional


def classificar_escolaridade(texto: str) -> str:
    niveis = set(re.findall(r"N[SIA]", texto.upper()))
    if niveis == {"NS"}:
        return "Superior"
    if niveis == {"NI"}:
        return "Médio"
    if niveis == {"NS", "NI"}:
        return "Superior e médio"
    if not niveis:
        return "Não informado"
    return "Inclui fundamental"


def extrair_data(*textos: str) -> date | None:
    """Procura '12 de dezembro de 2017', '11. Março. 2016' ou '-de-13-de-fevereiro-de-2025' (links)."""
    padrao = re.compile(r"(\d{1,2})\s*[º°o]?[\s.\-]*(?:de[\s\-]+)?([a-z]+)[\s.\-]*(?:de[\s\-]+)?(\d{4})")
    for texto in textos:
        for dia, mes, ano in padrao.findall(sem_acento(texto)):
            if mes in MESES:
                try:
                    return date(int(ano), MESES[mes], int(dia))
                except ValueError:
                    continue
    return None


def converter_vagas(valor: object) -> tuple[int | None, bool]:
    texto = limpar_texto(valor).upper()
    if texto in {"CR", "CADASTRO RESERVA"}:
        return 0, True
    if isinstance(valor, (int, float)):
        return int(valor), False
    digitos = re.sub(r"[^\d]", "", texto)
    return (int(digitos), False) if digitos else (None, False)


def ler_planilha(ano: int, caminho) -> list[dict]:
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    linhas = list(wb.worksheets[0].iter_rows(values_only=True))  # 1ª aba = tabela detalhada
    idx_cab = next(i for i, l in enumerate(linhas) if any(sem_acento(limpar_texto(c)).startswith("vagas") for c in l))
    col = identificar_colunas(linhas[idx_cab])
    pegar = lambda linha, nome: limpar_texto(linha[col[nome]]) if nome in col and col[nome] < len(linha) else ""

    registros, vistos = [], set()
    for linha in linhas[idx_cab + 1:]:
        orgao = pegar(linha, "orgao")
        vagas, cadastro_reserva = converter_vagas(linha[col["vagas"]] if col["vagas"] < len(linha) else None)
        if not orgao or sem_acento(orgao).startswith("total") or vagas is None:
            continue
        chave = tuple(limpar_texto(c) for c in linha)
        if chave in vistos:  # duplicata exata
            continue
        vistos.add(chave)

        bruto_cargos = str(linha[col["cargos"]] or "")
        cargos = [limpar_texto(c).strip("'") for c in re.split(r"[\n;|]", bruto_cargos) if limpar_texto(c)]
        ato, link, tipo_original = pegar(linha, "ato"), pegar(linha, "link"), pegar(linha, "tipo")
        data_ato = extrair_data(ato, link)
        registros.append({
            "ano": ano,
            "data": data_ato.isoformat() if data_ato else None,
            "orgao": orgao,
            "vinculo": pegar(linha, "vinculo") or orgao,
            "cargos": cargos,
            "escolaridade": classificar_escolaridade(pegar(linha, "escolaridade")),
            "vagas": vagas,
            "cadastroReserva": cadastro_reserva,
            "tipo": classificar_tipo(tipo_original),
            "tipoOriginal": tipo_original,
            "area": pegar(linha, "area") or "Não informada",
            "ato": ato,
            "link": link if link.startswith("http") else "",
            "ti": bool(PADRAO_TI.search(sem_acento(" ".join(cargos)))),
        })
    return registros


def descobrir_planilhas() -> dict[int, str]:
    """Lê a página oficial e devolve {ano: url} das planilhas anuais."""
    temporario = RAIZ / "data" / "tmp" / "pagina_concursos.html"
    try:
        baixar(PAGINA, temporario, tentativas=2, timeout=60)
        pagina = temporario.read_text(encoding="utf-8", errors="replace")
    except RuntimeError as erro:
        print(f"  ! página oficial indisponível ({erro}); usando os links conhecidos")
        return {}
    encontradas: dict[int, str] = {}
    # o ano às vezes fica partido por tags ("202</span>6") ou logo depois do link
    for m in re.finditer(r'<a\b[^>]*href="([^"]+\.xlsx[^"]*)"[^>]*>(.*?)</a>(.{0,120})', pagina, re.S | re.I):
        texto = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", m.group(2) + m.group(3)))
        ano = re.search(r"Provimentos\s*-?\s*(20\d\d)", texto)
        if ano and int(ano.group(1)) >= ANO_INICIAL:
            encontradas.setdefault(int(ano.group(1)), html.unescape(m.group(1)))
    print(f"  {len(encontradas)} planilhas anuais encontradas na página oficial")
    return encontradas


def ano_do_titulo(caminho) -> int | None:
    """Ano escrito no título da planilha ("Tabela de Autorizações e Provimentos - 2024")."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    for linha in wb.worksheets[0].iter_rows(max_row=5, values_only=True):
        for celula in linha:
            achou = re.search(r"Provimentos\s*-\s*(20\d\d)", str(celula or ""))
            if achou:
                return int(achou.group(1))
    return None


def atualizar_planilhas() -> dict[int, object]:
    """Baixa planilhas novas ou alteradas e devolve {ano: caminho} das válidas."""
    fontes: dict[str, str] = json.loads(FONTES.read_text(encoding="utf-8")) if FONTES.exists() else {}
    links = {**PLANILHAS_PADRAO, **descobrir_planilhas()}
    ultimo_ano = max(links)
    validas = {}
    for ano, url in sorted(links.items()):
        destino = DIR_CONCURSOS / f"autorizacoes_{ano}.xlsx"
        # baixa se é novo, se o link mudou ou se é um dos 2 anos mais recentes (ainda recebem atualizações)
        if not destino.exists() or fontes.get(str(ano)) != url or ano >= ultimo_ano - 1:
            temporario = RAIZ / "data" / "tmp" / f"autorizacoes_{ano}.xlsx"
            try:
                baixar(url, temporario, tentativas=3)
                escrever_se_mudou(temporario.read_bytes(), destino)
                fontes[str(ano)] = url
            except RuntimeError as erro:
                print(f"  ! {ano}: falha no download ({erro}); mantendo a versão anterior")
        if not destino.exists():
            continue
        ano_titulo = ano_do_titulo(destino)
        if ano_titulo != ano:
            print(f"  ! {ano}: o link oficial abre a tabela de {ano_titulo}; ano tratado como sem planilha")
            continue
        validas[ano] = destino
    escrever_se_mudou(json.dumps(dict(sorted(fontes.items())), indent=1, ensure_ascii=False).encode("utf-8"), FONTES)
    return validas


def main() -> None:
    print("Concursos federais (MGI)")
    planilhas = atualizar_planilhas()
    registros = []
    for ano, caminho in planilhas.items():
        dados = ler_planilha(ano, caminho)
        print(f"  {ano}: {len(dados):>3} autorizações, {sum(r['vagas'] for r in dados):>8,} vagas")
        registros += dados
    sem_planilha = [a for a in range(ANO_INICIAL, max(planilhas) + 1) if a not in planilhas]

    salvar_json(
        {
            "fonte": {
                "nome": "Tabelas de Autorizações e Provimentos — Ministério da Gestão e da Inovação (MGI)",
                "url": PAGINA,
                "geradoEm": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "anosSemPlanilha": sem_planilha,
                "ultimaDataAto": max(r["data"] for r in registros if r["data"]),
            },
            "registros": registros,
        },
        DIR_PUBLICO / "concursos.json",
    )


if __name__ == "__main__":
    main()
